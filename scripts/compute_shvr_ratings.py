"""
Parse SHVR (Swachh Evam Harit Vidyalaya Rating) per-district Excel exports for
Rajasthan and assign district-average star ratings onto hexes.

Source: data/raw/shvr_rajasthan/*.xlsx — one file per (post-2023-split)
district, exported from https://shvr.education.gov.in by a State Nodal
Officer account. Columns: S.No, School Name, UDISE Code, Address, School
Type, NEP Category, State Name, District Name, District Evaluation Status,
Self Evaluation Rating (1-5), Self Evaluation percentage.

No lat/lon in the source — schools are joined to hexes by DISTRICT, same
pattern as every other district-level dataset already in this platform
(NFHS-5, groundwater). The Excel files use Rajasthan's NEW (post-2023-split)
district names; the hex grid still uses the OLD 33 undivided districts, so
new districts are mapped back to their old parent (e.g. Balotra -> Barmer).

Coverage caveat (real, not a bug): only 32 of Rajasthan's ~50 current
districts were exported, covering 26 of the hex grid's 33 old districts.
7 old districts have zero data (Dungarpur, Ganganagar, Hanumangarh,
Jhunjhunun, Karauli, Rajsamand, Sawai Madhopur) and are left null, not
imputed. Includes ALL rows regardless of District Evaluation Status
(completed/not_assigned/in_progress/yet_to_start) per explicit instruction —
status is preserved per-school and per-district so the frontend can be
honest about which ratings reflect a finished evaluation.

Run: python scripts/compute_shvr_ratings.py
"""
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT      = Path(__file__).resolve().parent.parent
SRC_DIR   = ROOT / "data/raw/shvr_rajasthan"
HEX_PROPS = ROOT / "client/public/data/india_hex_props.json"

OUT_SCHOOLS  = ROOT / "client/public/data/shvr_schools_rajasthan.json"
OUT_DISTRICT = ROOT / "client/public/data/shvr_district_summary_rajasthan.json"
OUT_HEX      = ROOT / "client/public/data/shvr_hex_ratings_rajasthan.json"

# New (post-2023-split) district name, as it appears in the Excel files'
# "District Name" column -> old/undivided district name used by the hex grid.
# Districts not listed here match directly after .title() (e.g. AJMER -> Ajmer).
NEW_TO_OLD_DISTRICT = {
    "BALOTARA": "Barmer",
    "BEAWAR": "Ajmer",
    "DEEG": "Bharatpur",
    "DIDWANA-KUCHAMAN": "Nagaur",
    "KHAIRTHAL-TIJARA": "Alwar",
    "KOTPUTLI-BEHROR": "Jaipur",
    "PHALODI": "Jodhpur",
    "PRATAPGARH (RAJ.)": "Pratapgarh",
}


def map_district(raw: str) -> str:
    raw = (raw or "").strip().upper()
    return NEW_TO_OLD_DISTRICT.get(raw, raw.title())


def main():
    files = sorted(SRC_DIR.glob("*.xlsx"))
    print(f"Found {len(files)} SHVR district files")

    schools: list[dict] = []
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[1]:
                continue
            schools.append({
                "name": row[1],
                "udise_code": str(row[2]) if row[2] else None,
                "address": row[3],
                "school_type": row[4],
                "nep_category": row[5],
                "district_raw": row[7],
                "district": map_district(row[7]),
                "status": row[8],
                "rating": row[9],
                "percentage": float(row[10]) if row[10] not in (None, "") else None,
            })
    print(f"  {len(schools)} school records parsed")

    status_totals: dict[str, int] = defaultdict(int)
    for s in schools:
        status_totals[s["status"]] += 1
    print("  status breakdown:", dict(status_totals))

    # ── Per-district aggregate ──────────────────────────────────────────────
    by_district: dict[str, list[dict]] = defaultdict(list)
    for s in schools:
        by_district[s["district"]].append(s)

    district_summary: dict[str, dict] = {}
    for district, rows in by_district.items():
        rated = [r for r in rows if r["rating"] is not None]
        completed = [r for r in rows if r["status"] == "completed"]
        completed_rated = [r for r in completed if r["rating"] is not None]
        district_summary[district] = {
            "school_count": len(rows),
            "completed_count": len(completed),
            "pct_completed": round(100 * len(completed) / len(rows), 1) if rows else None,
            "avg_rating_all": round(sum(r["rating"] for r in rated) / len(rated), 2) if rated else None,
            "avg_rating_completed_only": round(sum(r["rating"] for r in completed_rated) / len(completed_rated), 2) if completed_rated else None,
            "rating_distribution_all": {str(k): sum(1 for r in rated if r["rating"] == k) for k in range(1, 6)},
            "status_counts": {k: sum(1 for r in rows if r["status"] == k) for k in set(r["status"] for r in rows)},
        }

    print(f"\n  {len(district_summary)} districts covered (mapped to old/hex district names):")
    for d, v in sorted(district_summary.items()):
        print(f"    {d:16s} n={v['school_count']:5d}  avg(all)={v['avg_rating_all']}  "
              f"avg(completed)={v['avg_rating_completed_only']}  completed={v['pct_completed']}%")

    # ── Join onto Rajasthan hexes (district-level, same as NFHS-5/groundwater) ──
    print(f"\nLoading {HEX_PROPS}...")
    props = json.loads(HEX_PROPS.read_text())
    raj_hexes = [p for p in props if p.get("state") == "Rajasthan"]
    print(f"  {len(raj_hexes)} Rajasthan hexes")

    hex_districts = sorted({p["district_name"] for p in raj_hexes if p.get("district_name")})
    covered = set(district_summary.keys()) & set(hex_districts)
    missing = sorted(set(hex_districts) - set(district_summary.keys()))
    print(f"  {len(covered)}/{len(hex_districts)} hex districts have SHVR data")
    print(f"  Missing entirely: {missing}")

    hex_ratings: dict[str, dict] = {}
    for p in raj_hexes:
        d = p.get("district_name")
        if d in district_summary:
            hex_ratings[p["h3_id"]] = {
                "district": d,
                "avg_rating_all": district_summary[d]["avg_rating_all"],
                "avg_rating_completed_only": district_summary[d]["avg_rating_completed_only"],
                "school_count": district_summary[d]["school_count"],
                "pct_completed": district_summary[d]["pct_completed"],
            }

    # ── Save ─────────────────────────────────────────────────────────────────
    print(f"\nSaving {OUT_SCHOOLS}...")
    OUT_SCHOOLS.write_text(json.dumps(schools, separators=(",", ":")))

    print(f"Saving {OUT_DISTRICT}...")
    OUT_DISTRICT.write_text(json.dumps({
        "meta": {
            "source": "SHVR (shvr.education.gov.in) per-district exports, State Nodal Officer access",
            "note": "32 of ~50 current Rajasthan districts exported, covering 26 of the hex grid's 33 "
                    "old/undivided districts. Includes ALL rows regardless of evaluation status — "
                    "avg_rating_completed_only is the more trustworthy figure; avg_rating_all also "
                    "includes not_assigned/in_progress/yet_to_start rows whose rating values may be "
                    "stale placeholders rather than finished evaluations.",
            "missing_districts": missing,
        },
        "districts": district_summary,
    }, indent=2))

    print(f"Saving {OUT_HEX}...")
    OUT_HEX.write_text(json.dumps(hex_ratings, separators=(",", ":")))

    import os
    print(f"\nDone.")
    print(f"  {OUT_SCHOOLS.name}: {os.path.getsize(OUT_SCHOOLS)//1024}KB")
    print(f"  {OUT_HEX.name}: {os.path.getsize(OUT_HEX)//1024}KB, {len(hex_ratings)} hexes rated")


if __name__ == "__main__":
    main()
